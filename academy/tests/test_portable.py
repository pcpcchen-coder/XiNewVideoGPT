"""Acceptance tests for the portable handoff; no TTS network or real L004 mutation."""
import sys
import shutil
import tempfile
import unittest
from unittest.mock import patch
import zipfile
from pathlib import Path
ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'pipeline'))
import academy as a
import portable_deck as deck

class PortableTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory()
        self.ep=Path(self.tmp.name)/'l004-portable-test'
        shutil.copytree(ROOT/'episodes/academy/l003-keyboard-ninja/production',self.ep/'production')
        m=a.read(self.ep/'production/manifest.json'); m.update(episode='L004',deckRenderer='ooxml-template',title='測試用可編輯投影片')
        a.write(self.ep/'production/manifest.json',m)
    def tearDown(self): self.tmp.cleanup()
    def test_template_edit_preserves_images_and_adds_current_notes(self):
        output=deck.build(self.ep,render=False)
        template=ROOT/'episodes/academy/l001-mac-workstation/production/presentation/l001-mac-workstation.pptx'
        with zipfile.ZipFile(output) as actual, zipfile.ZipFile(template) as source:
            media=[n for n in source.namelist() if n.startswith('ppt/media/')]
            self.assertTrue(media)
            for name in media: self.assertEqual(actual.read(name),source.read(name))
        from pptx import Presentation
        p=Presentation(output)
        self.assertEqual(len(p.slides),12)
        alltext='\n'.join(sh.text for s in p.slides for sh in s.shapes if sh.has_text_frame)
        self.assertNotIn('L001',alltext)
        self.assertIn('L004',alltext)
        self.assertIn(a.read(self.ep/'production/narration.json')[0]['sentences'][0]['text'],p.slides[0].notes_slide.notes_text_frame.text)
    def test_unmatched_mapping_rejected(self):
        edits=a.read(self.ep/'production/template-text.json'); edits[0]['old']='不存在的舊文字'
        a.write(self.ep/'production/template-text.json',edits)
        with self.assertRaisesRegex(ValueError,'Unmatched'): deck.build(self.ep,render=False)
    def test_paragraph_change_rejected(self):
        edits=a.read(self.ep/'production/template-text.json'); edits[0]['new']='一\n二\n三'
        a.write(self.ep/'production/template-text.json',edits)
        with self.assertRaisesRegex(ValueError,'paragraph count'): deck.build(self.ep,render=False)
    def test_portable_placeholder_cannot_bypass_validation(self):
        edits=a.read(self.ep/'production/template-text.json'); edits[0]['new']='待編寫'
        a.write(self.ep/'production/template-text.json',edits)
        with self.assertRaisesRegex(ValueError,'placeholders'): a.validate_episode(self.ep)
    def test_package_rejects_missing_delivery_even_with_technical_pass(self):
        (self.ep/'qc').mkdir(exist_ok=True)
        a.write(self.ep/'qc/report.json',{'checks':[{'status':'PASS'}]})
        a.write(self.ep/'qc/verified-inputs.json',a.fingerprint(self.ep))
        dest=Path(self.tmp.name)/'delivery'
        with patch.object(a,'media_check'), self.assertRaisesRegex(ValueError,'Incomplete delivery'):
            a.package(self.ep,dest)
        self.assertFalse(dest.exists())
    def test_classroom_change_invalidates_verified_fingerprint(self):
        f=self.ep/'classroom/task.md';f.parent.mkdir();f.write_text('v1')
        old=a.fingerprint(self.ep);f.write_text('v2')
        self.assertNotEqual(old,a.fingerprint(self.ep))
    def test_autospace_patch_rewrites_odp_default_and_rejects_missing_setting(self):
        tmp=Path(self.tmp.name); src=tmp/'in.odp'; dst=tmp/'out.odp'
        with zipfile.ZipFile(src,'w') as z:
            z.writestr('mimetype','application/vnd.oasis.opendocument.presentation')
            z.writestr('styles.xml','<s style:text-autospace="ideograph-alpha"/>')
            z.writestr('content.xml','<c/>')
        self.assertEqual(a.disable_asian_latin_autospace(src,dst),1)
        with zipfile.ZipFile(dst) as z:
            self.assertEqual(z.namelist()[0],'mimetype')
            self.assertIn('text-autospace="none"',z.read('styles.xml').decode())
        with zipfile.ZipFile(src,'w') as z:
            z.writestr('mimetype','x'); z.writestr('styles.xml','<s/>')
        with self.assertRaisesRegex(ValueError,'text-autospace'): a.disable_asian_latin_autospace(src,tmp/'no.odp')
    def test_caption_wrap_keeps_short_cues_and_breaks_long_ones_after_punctuation(self):
        sys.path.insert(0,str(ROOT/'pipeline/scripts'))
        import caption_wrap as cw
        short='引號很適合找一句話的出處、歌名或錯誤訊息；記得照畫面上的寫法，使用半形的雙引號。'
        self.assertEqual(cw.wrap_caption(short),[short])
        long='今天我們來玩搜尋尋寶：寶藏是一個大問題的答案，要找到它，得先把大問題拆成幾個好找的小問題。'
        lines=cw.wrap_caption(long)
        self.assertEqual(len(lines),2); self.assertEqual(''.join(lines),long)
        self.assertEqual(lines[0][-1],'，'); self.assertNotIn(lines[1][0],cw.NO_LINE_START)
        self.assertTrue(all(cw.width_em(x)<=cw.LINE_LIMIT_EM for x in lines))
        listy='第二種是精準搜尋，輸入颱風、防災、準備、清單這幾個關鍵字，再數一次，比較兩次結果差在哪裡。'
        self.assertEqual(cw.wrap_caption(listy)[0][-1],'，')  # sentence comma preferred over the list comma
        latin=cw.wrap_caption('一二三四五六七八九十一二三四五六七八九十一二三四五六七八九十一二三四 after:2026/01/01 一二三四五六')
        self.assertTrue(any('after:2026/01/01' in x for x in latin))  # never split inside a Latin/number run
        with self.assertRaisesRegex(ValueError,'two lines'): cw.wrap_caption('字'*90)
    def test_status_blocks_summarise_delivered_range_and_next_lesson(self):
        import status as st
        self.assertEqual(st._ranges(['L004','L006','L005','L009']),'L004–L006、L009')
        data=st.records(); blocks=st.status_blocks(data)
        nxt=next(x for x in data if x['productionStatus'] not in {'verified','completed_historical'})
        self.assertIn(nxt['lessonId'],blocks['sentence']); self.assertIn('不重製',blocks['sentence'])
        self.assertIn('大型輸出',blocks['large-files']); self.assertTrue(blocks['list'].startswith('\n- L'))
        for name in st.STATUS_DOCS: self.assertRegex((st.ROOT/name).read_text(encoding='utf-8'),st.MARK)
if __name__=='__main__': unittest.main()
