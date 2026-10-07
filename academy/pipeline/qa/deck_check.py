"""Reload the final PPTX and check editability, template leftovers and notes.

    ./portable-runtime.sh python pipeline/qa/deck_check.py --episode episodes/academy/<slug>

Structural only: it does not replace page-by-page viewing, and the deck is not
opened in PowerPoint or Keynote here.
"""
from pathlib import Path
import argparse
import json
import sys
import zipfile

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import academy as a  # noqa: E402

ALLOWED_SAME = {'→', '01', '02', '03', '00–25', '25–65', '65–75', '75–105', '105–115', '115–120',
                '休息十分鐘', '保持好奇，我們下次見！'}


def main():
    p = argparse.ArgumentParser(description=__doc__); p.add_argument('--episode', required=True, type=Path)
    ep = (a.ROOT / p.parse_args().episode).resolve()
    from pptx import Presentation
    m = a.read(ep / 'production/manifest.json')
    deck = ep / 'production/presentation' / f'{ep.name}.pptx'
    tpl = a.ROOT / m['templateEpisode'] / 'production/presentation' / (Path(m['templateEpisode']).name + '.pptx')
    prs = Presentation(deck); n = a.read(ep / 'production/narration.json'); edits = a.read(ep / 'production/template-text.json')
    texts = [[sh.text_frame.text for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip()] for s in prs.slides]
    alltext = '\n'.join(t for s in texts for t in s)
    with zipfile.ZipFile(deck) as z1, zipfile.ZipFile(tpl) as z0:
        media = [x for x in z0.namelist() if x.startswith('ppt/media/')]
        same = all(z1.read(x) == z0.read(x) for x in media)
    res = {
        'pptx': deck.name, 'bytes': deck.stat().st_size, 'sha256': a.sha(deck), 'slides': len(prs.slides),
        'nativeTextShapesPerSlide': [len(t) for t in texts],
        'picturesPerSlide': [sum(1 for sh in s.shapes if sh.shape_type == 13) for s in prs.slides],
        'templateMediaFiles': len(media), 'templateMediaBinaryIdentical': same,
        'allMappedTextPresent': all(e['new'] in texts[e['slide'] - 1] for e in edits),
        'leftoverTemplateTeachingText': [(e['slide'], e['old']) for e in edits
                                         if e['old'] not in ALLOWED_SAME and e['old'] in texts[e['slide'] - 1]],
        'containsL001Label': 'L001' in alltext, 'containsMacBookTitle': '把 MacBook' in alltext,
        'pageLabels': [next(t for t in s if t.startswith(m['episode'])) for s in texts],
        'notesContainAll36Sentences': all(all(x['text'] in s.notes_slide.notes_text_frame.text for x in sc['sentences'])
                                          for s, sc in zip(prs.slides, n)),
        'method': 'python-pptx reload of the final PPTX; binary comparison of ppt/media with the tracked template',
        'limits': 'not opened in Microsoft PowerPoint or Keynote; rendering reviewed via LibreOffice PNG export only'}
    res['passed'] = (res['slides'] == 12 and same and res['allMappedTextPresent'] and not res['leftoverTemplateTeachingText']
                     and not res['containsL001Label'] and not res['containsMacBookTitle'] and res['notesContainAll36Sentences'])
    a.write(ep / 'qc/editable-deck-check.json', res)
    print(json.dumps(res, ensure_ascii=False, indent=1))
    if not res['passed']:
        raise SystemExit(1)


if __name__ == '__main__':
    main()
