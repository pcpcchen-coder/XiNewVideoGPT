#!/usr/bin/env python3
"""Build the EP10 16:9 deck with the exact XiNewVideoGPT layout system.

Same chrome/layout geometry as EP01/EP03: theme-inherited fonts (no explicit
typeface), navy gradient bg, orange/blue chrome, rhino IP poses per slide.
"""
from __future__ import annotations

import argparse
import tempfile
import json
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from lxml import etree

ROOT = Path(__file__).resolve().parents[2]
parser = argparse.ArgumentParser()
parser.add_argument("--episode", required=True)
args = parser.parse_args()
EP = (ROOT / args.episode).resolve()
VIS = EP / "assets/visuals"
CUT = ROOT / "assets/characters/cutouts"

ORANGE = "FFC12F"; BLUE = "1982C4"; CYAN = "46C8FF"; WHITE = "F7F7F8"
GRAY = "A8B8C9"; PANEL = "101B28"; NAVY = "041222"; NAVY2 = "071F37"

EMU_IN = 914400
def IN(v): return Emu(int(round(v * EMU_IN)))

prs = Presentation()
prs.slide_width = Emu(12192000)
prs.slide_height = Emu(6858000)
BLANK = prs.slide_layouts[6]

manifest = json.loads((EP / "production/manifest.json").read_text(encoding="utf-8"))
slides_cfg = json.loads((EP / "production/slides.json").read_text(encoding="utf-8"))
narration = json.loads((EP / "production/narration.json").read_text(encoding="utf-8"))
TITLE = manifest["title"]; EPISODE = manifest["episode"].replace("S01E", "EP")
COUNT = manifest["slideCount"]

def set_font(run, size, bold, color):
    # match EP01/EP03: no explicit typeface -> inherits theme, system CJK fallback
    f = run.font
    f.size = Pt(size); f.bold = bold
    f.color.rgb = RGBColor.from_string(color)

def bg_gradient(slide):
    xml = (
        '<p:bg xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
        'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><p:bgPr>'
        '<a:gradFill><a:gsLst>'
        f'<a:gs pos="0"><a:srgbClr val="{NAVY}"/></a:gs>'
        f'<a:gs pos="62000"><a:srgbClr val="{NAVY2}"/></a:gs>'
        f'<a:gs pos="100000"><a:srgbClr val="{NAVY}"/></a:gs>'
        '</a:gsLst><a:lin ang="8100000"/></a:gradFill>'
        '<a:effectLst/></p:bgPr></p:bg>')
    bg = etree.fromstring(xml)
    cSld = slide._element.find(qn("p:cSld"))
    cSld.insert(0, bg)

def rect(slide, x, y, w, h, fill=None, alpha=None, line=None, line_w=1.0, round_=False):
    shp = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if round_ else MSO_SHAPE.RECTANGLE,
        IN(x), IN(y), IN(w), IN(h))
    shp.shadow.inherit = False
    if round_:
        try: shp.adjustments[0] = 0.16
        except Exception: pass
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid(); shp.fill.fore_color.rgb = RGBColor.from_string(fill)
        if alpha is not None:
            srgb = shp.fill.fore_color._xFill.find(qn("a:srgbClr"))
            a = srgb.makeelement(qn("a:alpha"), {"val": str(int(alpha * 1000))})
            srgb.append(a)
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = RGBColor.from_string(line); shp.line.width = Pt(line_w)
    return shp

def text(slide, x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, line_spacing=None):
    tb = slide.shapes.add_textbox(IN(x), IN(y), IN(w), IN(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    first = True
    for line in runs:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = align
        if line_spacing: p.line_spacing = line_spacing
        for (t, size, bold, color) in line:
            r = p.add_run(); r.text = t
            set_font(r, size, bold, color)
    return tb

def pic(slide, path, x, y, w, h, crop=None):
    p = slide.shapes.add_picture(str(path), IN(x), IN(y), IN(w), IN(h))
    if crop:
        p.crop_left, p.crop_right = crop.get("l", 0), crop.get("r", 0)
        p.crop_top, p.crop_bottom = crop.get("t", 0), crop.get("b", 0)
    p.shadow.inherit = False
    return p

POSE_FILE = {"point-up": "xinew-point-up.png", "pointer": "xinew-pointer.png",
             "open-palms": "xinew-open-palms.png", "thumbs-up": "xinew-thumbs-up.png"}
def rhino(slide, pose, x, y, w=3.23):
    return pic(slide, CUT / POSE_FILE[pose], x, y, w, w / 1.777)

def chrome(slide, n, eyebrow):
    rect(slide, 0, 0, 13.333, 0.0833, fill=ORANGE)
    rect(slide, 0.58, 0.90, 12.17, 0.0208, fill=BLUE, alpha=55)
    text(slide, 0.58, 0.31, 1.09, 0.38, [[("CHEN", 18.75, True, WHITE)]])
    text(slide, 1.56, 0.31, 0.99, 0.38, [[("XiNew", 18.75, True, ORANGE)]])
    text(slide, 2.81, 0.38, 5.42, 0.29, [[(eyebrow, 11.25, True, GRAY)]])
    text(slide, 10.99, 0.35, 1.77, 0.31, [[(f"{EPISODE} · {n:02d}/{COUNT}", 12.0, True, ORANGE)]])
    text(slide, 0.58, 7.04, 5.21, 0.25, [[(TITLE, 9.75, False, GRAY)]])
    text(slide, 11.30, 7.04, 1.46, 0.25, [[("XINEW SAYS", 9.75, True, ORANGE)]])

def new_slide(n):
    s = prs.slides.add_slide(BLANK)
    bg_gradient(s)
    return s

def two_columns(s, cfg, pose):
    """Windows vs macOS compare layout (EP03 slide 9 geometry)."""
    text(s, 0.60, 1.17, 11.5, 0.81, [[(cfg["title"], 36.0, True, WHITE)]])
    for side, (tt, items, xc, col) in enumerate([
            (cfg["leftTitle"], cfg["leftItems"], 0.65, CYAN),
            (cfg["rightTitle"], cfg["rightItems"], 7.12, ORANGE)]):
        rect(s, xc, 2.29, 5.52, 3.44, fill=PANEL, alpha=88, line=col, round_=True)
        text(s, xc + 0.29, 2.58, 4.90, 0.47, [[(tt, 21.0, True, col)]])
        for i, it in enumerate(items):
            text(s, xc + 0.37, 3.33 + i * 0.605, 4.58, 0.42, [[("•  " + it, 17.25, False, WHITE)]])
    rect(s, 2.97, 6.09, 7.40, 0.57, fill=ORANGE, alpha=14, line=ORANGE, round_=True)
    text(s, 3.18, 6.21, 6.98, 0.39, [[(cfg["takeaway"], 15.0, True, ORANGE)]])
    rhino(s, pose, 5.72, 4.98, 1.90)

def four_cards(s, cfg, pose, visual):
    """Full-width visual band + 4 cards (EP03 slide 10 geometry)."""
    text(s, 0.60, 1.12, 12.4, 0.81, [[(cfg["title"], 34.5, True, WHITE)]])
    pic(s, VIS / f"{visual}.png", 0.57, 2.19, 12.19, 4.06, crop={"t": 0.20386, "b": 0.20386})
    hi = 2
    for i, col_cfg in enumerate(cfg["columns"]):
        x = 0.94 + i * 2.92
        hot = (i == hi)
        col = ORANGE if hot else CYAN
        rect(s, x, 5.05, 2.72, 1.00, fill=NAVY, alpha=88, line=ORANGE if hot else BLUE, round_=True)
        text(s, x + 0.21, 5.21, 2.30, 0.35, [[(col_cfg["title"], 18.0, True, col)]])
        text(s, x + 0.21, 5.60, 2.30, 0.29, [[(col_cfg["subtitle"], 13.5, False, WHITE)]])
    rhino(s, pose, 10.90, 3.80, 2.20)

poses = manifest["poses"]
for cfg in slides_cfg:
    n = cfg["slide"]; pose = poses[n - 1]
    s = new_slide(n)
    chrome(s, n, cfg["eyebrow"])

    if n == 1:  # cover
        pic(s, VIS / "future-hero.png", 0, 0.0833, 13.333, 6.958, crop={"t": 0.03636, "b": 0.03636})
        rect(s, 0, 0.0833, 13.333, 0.0833, fill=ORANGE)
        text(s, 0.60, 0.81, 4.79, 0.31, [[(cfg["eyebrow"], 13.5, True, ORANGE)]])
        text(s, 0.60, 1.35, 6.4, 1.88, [[(line, 52.5, True, WHITE)] for line in cfg["title"].split("\n")])
        text(s, 0.65, 3.44, 6.6, 0.54, [[(cfg["subtitle"], 21.0, False, WHITE)]])
        rect(s, 0.62, 4.38, 5.68, 0.60, fill=ORANGE, alpha=14, line=ORANGE, round_=True)
        text(s, 0.83, 4.49, 5.26, 0.42, [[(cfg["takeaway"], 15.0, True, ORANGE)]])
        rhino(s, pose, 4.27, 4.35, 3.23)

    elif n in (2, 5, 11):  # step flow (3 or 4 steps)
        text(s, 0.60, 1.17, 11.5, 0.81, [[(cfg["title"], 36.0, True, WHITE)]])
        text(s, 0.62, 1.93, 10.5, 0.42, [[(cfg["subtitle"], 16.5, False, GRAY)]])
        hi = {2: 1, 5: 2, 11: 2}.get(n, 1)
        three = len(cfg["steps"]) == 3
        cw, step, tw = (2.30, 2.75, 2.10) if three else (1.51, 1.973, 1.30)
        for i, st in enumerate(cfg["steps"]):
            x = 0.67 + i * step
            hot = (i == hi)
            rect(s, x, 3.33, cw, 0.96, fill=ORANGE if hot else BLUE, alpha=18 if hot else 20, line=BLUE, round_=True)
            text(s, x + 0.10, 3.46, tw, 0.77, [[(st, 19.5, True, ORANGE if hot else WHITE)]])
            if i < len(cfg["steps"]) - 1:
                text(s, x + cw + 0.03, 3.48, 0.44, 0.62, [[("→", 28.5, True, CYAN)]])
        rect(s, 0.65, 4.95, 9.45, 0.83, fill=BLUE, alpha=10, line=BLUE, round_=True)
        text(s, 0.96, 5.21, 8.85, 0.40, [[(cfg["takeaway"], 18.0, False, WHITE)]])
        rhino(s, pose, 10.10, 3.44, 3.23)

    elif n in (3, 8):  # full-bleed analogy / story
        vis = {3: "day-map", 8: "tutor-map"}[n]
        pic(s, VIS / f"{vis}.png", 0, 0.92, 13.333, 6.115, crop={"t": 0.09188, "b": 0.09188})
        text(s, 0.60, 1.15, 11.0, 0.81, [[(cfg["title"], 36.0, True, WHITE)]])
        text(s, 0.62, 1.94, 11.56, 0.42, [[(cfg["subtitle"], 16.5, False, WHITE)]])
        hi = 2
        for i, lab in enumerate(cfg["labels"]):
            x = 0.56 + i * 2.4075
            rect(s, x, 6.08, 2.08, 0.56, fill=NAVY, alpha=80, line=BLUE, round_=True)
            text(s, x + 0.11, 6.21, 1.88, 0.38, [[(lab, 12.75, True, ORANGE if i == hi else WHITE)]])
        if n == 3:
            rhino(s, pose, 10.90, 3.70, 2.20)
        else:
            rhino(s, pose, 11.05, 5.00, 2.00)

    elif n in (6, 9):  # two columns compare
        two_columns(s, cfg, pose)

    elif n == 10:  # full-width visual + 4 cards
        four_cards(s, cfg, pose, "drive-wheel")

    elif n in (4, 7):  # left visual + right bullets
        vis = {4: "agent-tasks", 7: "creative-director"}[n]
        text(s, 0.60, 1.12, 12.4, 0.81, [[(cfg["title"], 34.5, True, WHITE)]])
        pic(s, VIS / f"{vis}.png", 0.60, 2.27, 5.94, 3.75, crop={"l": 0.05445, "r": 0.05445})
        for i, b in enumerate(cfg["bullets"]):
            y = 2.79 + i * 0.835
            hot = (i == 2)
            dot = s.shapes.add_shape(MSO_SHAPE.OVAL, IN(7.40), IN(y), IN(0.21), IN(0.21))
            dot.shadow.inherit = False
            dot.fill.solid(); dot.fill.fore_color.rgb = RGBColor.from_string(ORANGE if hot else CYAN)
            dot.line.fill.background()
            text(s, 7.83, y - 0.19, 4.6, 0.57, [[(b, 19.5, hot, WHITE)]])
        rect(s, 7.29, 5.21, 5.21, 0.62, fill=ORANGE, alpha=14, line=ORANGE, round_=True)
        text(s, 7.50, 5.32, 4.79, 0.44, [[(cfg["takeaway"], 15.0, True, ORANGE)]])
        rhino(s, pose, 5.36, 4.66, 1.41)

    elif n == 12:  # quiz
        text(s, 0.60, 1.12, 11.5, 0.81, [[(cfg["title"], 34.5, True, WHITE)]])
        rect(s, 0.65, 2.29, 7.29, 2.34, fill=PANEL, alpha=92, line=BLUE, round_=True)
        text(s, 0.98, 2.58, 0.94, 0.35, [[("答案", 13.5, True, GRAY)]])
        text(s, 0.96, 3.04, 4.79, 0.73, [[(cfg["answer"], 36.0, True, ORANGE)]])
        text(s, 0.98, 3.91, 6.35, 0.44, [[(cfg["explain"], 17.25, False, WHITE)]])
        hi = 1
        for i, rc in enumerate(cfg["recap"]):
            x = 0.69 + i * 1.465
            hot = (i == hi)
            rect(s, x, 5.00, 1.25, 0.60, fill=ORANGE if hot else BLUE, alpha=18, line=BLUE, round_=True)
            text(s, x + 0.10, 5.12, 1.04, 0.42, [[(rc, 15.0, True, ORANGE if hot else WHITE)]])
        text(s, 0.73, 6.09, 6.0, 0.46, [[(cfg["next"], 18.75, True, CYAN)]])
        rhino(s, pose, 8.23, 2.89, 4.90)
        if cfg.get("signoff"):
            rect(s, 6.99, 6.02, 5.75, 0.62, fill=ORANGE, alpha=14, line=ORANGE, round_=True)
            text(s, 7.10, 6.13, 5.53, 0.42, [[(cfg["signoff"], 15.0, True, ORANGE)]], align=PP_ALIGN.CENTER)

    # speaker notes = narration sentences
    notes = prs.slides[n - 1].notes_slide
    tf = notes.notes_text_frame
    sc = narration[n - 1]
    tf.text = f"第 {n} 幕：{sc['title']}"
    for sent in sc["sentences"]:
        p = tf.add_paragraph(); p.text = sent["text"]

out = EP / "production/presentation" / f"{EP.name}.pptx"
out.parent.mkdir(parents=True, exist_ok=True)
import shutil, tempfile, time
tmp_dir = tempfile.TemporaryDirectory(prefix="xinew-deck-")
tmp = Path(tmp_dir.name) / out.name
prs.save(tmp)
for attempt in range(5):
    try:
        shutil.copyfile(tmp, out)
        break
    except OSError:
        time.sleep(1)
print("saved", out)
